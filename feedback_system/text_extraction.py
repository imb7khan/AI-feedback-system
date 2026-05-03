"""
Text extraction service for processing uploaded files.

Supports PDF, DOCX, and plain text files with comprehensive error handling.
"""

import os
from typing import Optional
import logging

from PyPDF2 import PdfReader
from docx import Document

logger = logging.getLogger(__name__)


class TextExtractionError(Exception):
    """Custom exception for text extraction errors."""
    pass


class TextExtractor:
    """Service class for extracting text from various file formats."""

    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """
        Extract text from a file based on its type.

        Args:
            file_path: Path to the file to extract text from
            file_type: Type of file ('pdf', 'docx', 'txt')

        Returns:
            Extracted text as a string

        Raises:
            TextExtractionError: If extraction fails
            FileNotFoundError: If file doesn't exist
        """
        # Validate inputs
        if not file_path or not isinstance(file_path, str):
            raise TextExtractionError("Invalid file path provided")

        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        # Check file size (limit to 10MB)
        file_size = os.path.getsize(file_path)
        max_size = 10 * 1024 * 1024  # 10MB
        if file_size > max_size:
            raise TextExtractionError(
                f"File too large ({file_size / 1024 / 1024:.1f}MB). "
                f"Maximum size is {max_size / 1024 / 1024}MB."
            )

        if file_size == 0:
            raise TextExtractionError("File is empty")

        # Validate file type
        valid_types = ['pdf', 'docx', 'txt']
        if file_type not in valid_types:
            raise TextExtractionError(
                f"Unsupported file type: {file_type}. "
                f"Supported types: {', '.join(valid_types)}"
            )

        try:
            if file_type == 'pdf':
                return TextExtractor._extract_from_pdf(file_path)
            elif file_type == 'docx':
                return TextExtractor._extract_from_docx(file_path)
            elif file_type == 'txt':
                return TextExtractor._extract_from_txt(file_path)
            else:
                raise TextExtractionError(
                    f"Unsupported file type: {file_type}"
                )
        except TextExtractionError:
            # Re-raise TextExtractionError as-is
            raise
        except FileNotFoundError:
            # Re-raise FileNotFoundError as-is
            raise
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {str(e)}")
            raise TextExtractionError(f"Failed to extract text: {str(e)}")

    @staticmethod
    def _extract_from_pdf(file_path: str) -> str:
        """
        Extract text from a PDF file.

        Args:
            file_path: Path to the PDF file

        Returns:
            Extracted text as a string

        Raises:
            TextExtractionError: If PDF extraction fails
        """
        try:
            reader = PdfReader(file_path)

            # Check if PDF is encrypted
            if reader.is_encrypted:
                raise TextExtractionError(
                    "PDF is encrypted and password protected. "
                    "Please provide an unencrypted PDF."
                )

            # Check if PDF has any pages
            if len(reader.pages) == 0:
                raise TextExtractionError("PDF has no pages.")

            text_parts = []
            total_pages = len(reader.pages)

            for page_num, page in enumerate(reader.pages, 1):
                try:
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        text_parts.append(page_text)
                    else:
                        logger.warning(f"Page {page_num} of {total_pages} contains no extractable text (might be image-based)")
                except Exception as e:
                    logger.warning(f"Error extracting text from page {page_num}: {str(e)}")
                    continue

            extracted_text = '\n'.join(text_parts)

            if not extracted_text.strip():
                raise TextExtractionError(
                    "No text could be extracted from PDF. "
                    "The PDF might be scanned, image-based, or use unsupported encoding."
                )

            # Limit text length to prevent memory issues
            max_length = 100000  # 100k characters
            if len(extracted_text) > max_length:
                logger.warning(f"PDF text truncated from {len(extracted_text)} to {max_length} characters")
                extracted_text = extracted_text[:max_length] + "\n\n[Text truncated due to length]"

            return extracted_text

        except TextExtractionError:
            # Re-raise TextExtractionError as-is
            raise
        except Exception as e:
            raise TextExtractionError(f"PDF extraction failed: {str(e)}")

    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        """
        Extract text from a DOCX file.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Extracted text as a string

        Raises:
            TextExtractionError: If DOCX extraction fails
        """
        try:
            doc = Document(file_path)
            text_parts = []

            # Check if document has any paragraphs
            if not doc.paragraphs:
                raise TextExtractionError("DOCX file contains no paragraphs.")

            for paragraph in doc.paragraphs:
                if paragraph.text and paragraph.text.strip():
                    text_parts.append(paragraph.text)

            extracted_text = '\n'.join(text_parts)

            if not extracted_text.strip():
                raise TextExtractionError(
                    "No text could be extracted from DOCX file. "
                    "The file might be empty or contain only images."
                )

            # Limit text length to prevent memory issues
            max_length = 100000  # 100k characters
            if len(extracted_text) > max_length:
                logger.warning(f"DOCX text truncated from {len(extracted_text)} to {max_length} characters")
                extracted_text = extracted_text[:max_length] + "\n\n[Text truncated due to length]"

            return extracted_text

        except TextExtractionError:
            # Re-raise TextExtractionError as-is
            raise
        except Exception as e:
            raise TextExtractionError(f"DOCX extraction failed: {str(e)}")

    @staticmethod
    def _extract_from_txt(file_path: str) -> str:
        """
        Extract text from a plain text file.

        Args:
            file_path: Path to the text file

        Returns:
            Extracted text as a string

        Raises:
            TextExtractionError: If text file reading fails
        """
        try:
            # Try different encodings in order of preference
            encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'iso-8859-1']

            last_error = None
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        text = file.read()

                    if text and text.strip():
                        # Normalize line endings
                        text = text.replace('\r\n', '\n').replace('\r', '\n')

                        # Limit text length to prevent memory issues
                        max_length = 100000  # 100k characters
                        if len(text) > max_length:
                            logger.warning(f"Text file truncated from {len(text)} to {max_length} characters")
                            text = text[:max_length] + "\n\n[Text truncated due to length]"

                        return text

                except UnicodeDecodeError as e:
                    last_error = e
                    logger.debug(f"Failed to decode with {encoding}: {str(e)}")
                    continue
                except Exception as e:
                    last_error = e
                    logger.debug(f"Error reading with {encoding}: {str(e)}")
                    continue

            # If all encodings failed
            raise TextExtractionError(
                f"Could not decode text file with supported encodings. "
                f"Last error: {str(last_error) if last_error else 'Unknown encoding error'}"
            )

        except TextExtractionError:
            # Re-raise TextExtractionError as-is
            raise
        except Exception as e:
            raise TextExtractionError(f"Text file reading failed: {str(e)}")

    @staticmethod
    def extract_text_from_file_object(file_object, file_type: str) -> str:
        """
        Extract text from a Django file object.

        Args:
            file_object: Django FileField object
            file_type: Type of file ('pdf', 'docx', 'txt')

        Returns:
            Extracted text as a string

        Raises:
            TextExtractionError: If extraction fails
        """
        try:
            # Get the file path
            file_path = file_object.path

            # If file doesn't exist yet, save it first
            if not os.path.exists(file_path):
                file_object.save(file_path, file_object)

            return TextExtractor.extract_text(file_path, file_type)

        except Exception as e:
            raise TextExtractionError(f"Failed to extract text from file object: {str(e)}")